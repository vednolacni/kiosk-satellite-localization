#!/usr/bin/env python3
"""Build sl-review.xlsx for the native-speaker review.

Usage: review_xlsx.py REPO FLAGS_TSV OUT.xlsx

Sheet "Pregled": every message in source order. Sheet "Zastavice": only
flagged messages. Columns: datoteka, ključ, angleško, slovensko,
context/description, zastavica, opomba k zastavici, Popravek (empty).
"""
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["datoteka", "ključ", "angleško", "slovensko", "context/description",
           "zastavica", "opomba k zastavici", "Popravek"]
WIDTHS = [30, 38, 60, 60, 50, 22, 50, 40]
FILLS = {
    "negotovo": "FFF2CC",
    "preoblikovano zaradi števila": "DDEBF7",
    "placeholder v sklonu": "E2EFDA",
    "drugo": "EDEDED",
}


def meta_text(meta):
    parts = []
    if meta.get("context"):
        parts.append("Context: " + meta["context"])
    if meta.get("description"):
        parts.append("Description: " + meta["description"])
    if meta.get("x-notes"):
        parts.append("Notes: " + meta["x-notes"])
    for name, info in (meta.get("placeholders") or {}).items():
        if isinstance(info, dict) and info.get("description"):
            parts.append(f"{{{name}}}: {info['description']}")
    return "\n".join(parts)


def setup_sheet(ws):
    ws.append(HEADERS)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="BDD7EE")
    ws.freeze_panes = "C2"


def add_row(ws, row):
    ws.append(row)
    r = ws.max_row
    for i in range(1, len(HEADERS) + 1):
        ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    flag = row[5]
    if flag in FILLS:
        ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=FILLS[flag])


def main():
    repo, flags_path, out = sys.argv[1:4]
    flags = {}
    for line in open(flags_path, encoding="utf-8").read().splitlines()[1:]:
        parts = line.split("\t")
        if len(parts) >= 4:
            flags[(parts[0], parts[1])] = (parts[2], parts[3])

    wb = Workbook()
    all_ws = wb.active
    all_ws.title = "Pregled"
    flag_ws = wb.create_sheet("Zastavice")
    setup_sheet(all_ws)
    setup_sheet(flag_ws)

    count = flagged = 0
    for path in sorted(glob.glob(os.path.join(repo, "source", "*_en.arb"))):
        name = os.path.basename(path)[: -len("_en.arb")]
        src = json.load(open(path, encoding="utf-8"))
        sl = json.load(open(os.path.join(repo, "translations", "sl", name + "_sl.arb"), encoding="utf-8"))
        for key, en in src.items():
            if key.startswith("@"):
                continue
            flag, note = flags.get((name, key), ("", ""))
            row = [name, key, en, sl.get(key, ""), meta_text(src.get("@" + key, {})), flag, note, ""]
            add_row(all_ws, row)
            count += 1
            if flag:
                add_row(flag_ws, row)
                flagged += 1

    for ws in (all_ws, flag_ws):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    wb.save(out)
    print(f"wrote {out}: {count} messages, {flagged} flagged")


if __name__ == "__main__":
    main()
